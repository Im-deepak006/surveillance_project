from pathlib import Path
import json,textwrap
from PIL import Image,ImageDraw,ImageFont
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
R=Path(__file__).resolve().parent; O=R/'THESIS_DELIVERABLES_DRAFT_2026-10-24'; F=O/'figures'; F.mkdir(parents=True,exist_ok=False)
D=json.loads((R/'outputs/attribute_aware_consensus_20261023/results.json').read_text()); M=D['metrics']
cs=['A_temporal_only','B_saved_single_frame_appearance','D_saved_person_object_fusion','E_median_multiframe_appearance','F_median_appearance_plus_repeated_object','G_single_frame_appearance_plus_repeated_object']
labs=['A temporal','B single-frame','D Method D','E median 3-frame','F E + object','G B + object']
def ft(n,b=False):
 p=('C:/Windows/Fonts/arialbd.ttf' if b else 'C:/Windows/Fonts/arial.ttf')
 return ImageFont.truetype(p,n) if Path(p).exists() else ImageFont.load_default()
def diagram(name,heading,items):
 im=Image.new('RGB',(1400,440),'white'); d=ImageDraw.Draw(im); d.text((700,35),heading,font=ft(27,True),fill='#17324D',anchor='mm'); step=1320//len(items)
 for i,(a,b) in enumerate(items):
  x=40+i*step; d.rounded_rectangle((x,160,x+step-28,270),14,fill='#EAF1F8',outline='#2878A5',width=3); d.text((x+(step-28)/2,190),a,font=ft(18,True),fill='#17324D',anchor='mm'); d.multiline_text((x+(step-28)/2,230),textwrap.fill(b,22),font=ft(13),fill='#35556F',anchor='mm',align='center')
  if i<len(items)-1:
   d.line((x+step-28,215,x+step-12,215),fill='#536B82',width=4); d.polygon([(x+step-12,215),(x+step-24,208),(x+step-24,222)],fill='#536B82')
 im.save(F/name)
diagram('fig01_architecture.png','System architecture',[('Input','video + query'),('UI','Streamlit'),('Baseline','Moment-DETR'),('Evidence','YOLO + OpenCLIP'),('Fusion','rank mean'),('Review','FFmpeg clips')])
diagram('fig02_baseline.png','Baseline Moment-DETR pipeline',[('Video','2 s sampling'),('CLIP','512-D features'),('TEF','temporal endpoints'),('Transformer','2 enc + 2 dec'),('Output','10 intervals')])
diagram('fig03_method_d.png','Method D reranking',[('Proposals','fixed 10'),('Evidence','temporal / crop / object'),('Ranks','tied percentiles'),('Fusion','unweighted mean'),('Output','reordered intervals')])
diagram('fig04_dataflow.png','Video-to-ranked-moment data flow',[('Upload','video and text'),('Encode','CLIP features'),('Propose','Moment-DETR'),('Score','optional cues'),('Rank','normalize/fuse'),('Display','validate/play')])
def chart(name,title,key,mx):
 im=Image.new('RGB',(1200,590),'white');d=ImageDraw.Draw(im);d.text((600,30),title,font=ft(25,True),fill='#17324D',anchor='mm'); L,B,T=80,455,90;d.line((L,T,L,B),fill='#536B82',width=2);d.line((L,B,1170,B),fill='#536B82',width=2)
 for k in range(6):
  y=B-k*(B-T)/5;d.line((L,y,1170,y),fill='#D9E1E8');d.text((L-8,y),f'{mx*k/5:.1f}',font=ft(13),fill='#35556F',anchor='rm')
 group=1080/6
 for i,c in enumerate(cs):
  val=M[c][key]; x=L+i*group+24;y=B-(B-T)*val/mx;d.rectangle((x,y,x+70,B),fill='#D27A2C' if i==2 else '#2878A5');d.text((x+35,y-13),f'{val:.3f}',font=ft(14,True),fill='#17324D',anchor='mm');d.multiline_text((x+35,B+12),textwrap.fill(labs[i],12),font=ft(12),fill='#35556F',anchor='ma',align='center')
 im.save(F/name)
chart('fig05_mean_iou.png','Top-1 mean IoU (n=5)','mean_top1_iou',.5)
chart('fig06_recall.png','Recall@1 at IoU 0.5','recall_at_1_iou_0_5',1)
src=R/'outputs/automated_visual_diagnostics_mot17_04_verified_20261023/images'; imgs=sorted(src.glob('*_D_t*.jpg'))[:5]
if imgs:
 im=Image.new('RGB',(1100,100+350*((len(imgs)+1)//2)),'white');d=ImageDraw.Draw(im);d.text((550,30),'Actual saved MOT17-04 diagnostic frames',font=ft(23,True),fill='#17324D',anchor='mm')
 for i,p in enumerate(imgs):
  a=Image.open(p).convert('RGB');a.thumbnail((510,285));x=25+i%2*540;y=70+i//2*350;im.paste(a,(x,y));d.text((x,y+292),p.stem,font=ft(11),fill='#17324D')
 im.save(F/'fig07_actual_frames.png')
def setup(doc,title):
 s=doc.sections[0];s.top_margin=Inches(.72);s.bottom_margin=Inches(.72);s.left_margin=Inches(.8);s.right_margin=Inches(.8)
 for n,z in [('Normal',10.5),('Title',24),('Heading 1',17),('Heading 2',12)]:
  st=doc.styles[n];st.font.name='Cambria';st.font.size=Pt(z)
  if n!='Normal':st.font.bold=True;st.font.color.rgb=RGBColor(23,50,77)
 doc.styles['Normal'].paragraph_format.space_after=Pt(6);doc.styles['Normal'].paragraph_format.line_spacing=1.1
 p=s.footer.paragraphs[0];p.alignment=1;p.add_run('DRAFT • Project thesis | Page ');f=OxmlElement('w:fldSimple');f.set(qn('w:instr'),'PAGE');p._p.append(f);doc.core_properties.title=title
def pg(d):d.add_page_break()
def hd(d,s):d.add_heading(s,0)
def h(d,s):d.add_heading(s,2)
def p(d,s):d.add_paragraph(s)
def bullets(d,xs):
 for x in xs:d.add_paragraph(x,style='List Bullet')
def tbl(d,heads,rows):
 t=d.add_table(rows=1,cols=len(heads));t.style='Light Shading Accent 1'
 for i,x in enumerate(heads):t.rows[0].cells[i].text=x
 for row in rows:
  c=t.add_row().cells
  for i,x in enumerate(row):c[i].text=str(x)
 for row in t.rows:
  for c in row.cells:
   for q in c.paragraphs:
    for r in q.runs:r.font.size=Pt(8)
def pic(d,name,cap):
 q=d.add_paragraph();q.alignment=1;q.add_run().add_picture(str(F/name),width=Inches(6.4));q=d.add_paragraph(cap);q.alignment=1
t=Document();setup(t,'Attribute-Aware Video Moment Retrieval for Person-Centric Search')
q=t.add_paragraph();q.alignment=1;q.paragraph_format.space_before=Pt(55);r=q.add_run('ATTRIBUTE-AWARE VIDEO MOMENT\\nRETRIEVAL FOR PERSON-CENTRIC SEARCH');r.bold=True;r.font.size=Pt(25);r.font.color.rgb=RGBColor(23,50,77)
for x in ['Final-Year B.Tech Information Technology Project Thesis','[UNIVERSITY DEGREE WORDING]','Student: [NAME]  |  Register number: [NUMBER]','Guide: [SUPERVISOR NAME AND DESIGNATION]','[DEPARTMENT]  |  [INSTITUTION]  |  [CITY]','[SUBMISSION MONTH AND YEAR]']:
 q=t.add_paragraph(x);q.alignment=1;q.paragraph_format.space_before=Pt(16)
pg(t);hd(t,'Certificate');p(t,'[Insert the approved institutional certificate; obtain required signatures and seal. No approval is represented here.]')
pg(t);hd(t,'Declaration');p(t,'[Insert and sign the prescribed declaration after confirming authorship/originality wording. Add student details, date, and place.]')
pg(t);hd(t,'Acknowledgements');p(t,'We thank [GUIDE NAME] and [DEPARTMENT / INSTITUTION] for their guidance and academic support. We acknowledge the authors and maintainers of Moment-DETR, QVHighlights, CLIP, OpenCLIP, Ultralytics, PyTorch, OpenCV, and FFmpeg. [Add direct contributors and confirm wording.]')
pg(t);hd(t,'Abstract');p(t,'Manual review of long recordings is time-consuming when a reviewer must locate a short event from natural language. This project integrates a pretrained Moment-DETR model into a prototype that accepts a short video and query and returns ranked temporal intervals. The proposed Method D contribution is inference-time reranking: person-crop and available person-object/category evidence are converted into candidate ranks and fused with temporal rank. The checkpoint and architecture remain unchanged. Saved results cover five manually defined queries on one MOT17-04-SDP video. The temporal baseline records mean top-1 Intersection over Union (IoU) 0.1940 and Recall@1 of 0.4, 0.2, and 0.0 at thresholds 0.3, 0.5, and 0.7. Method D records 0.4020 mean IoU and recalls 0.8, 0.4, and 0.0, an absolute difference of 0.2080. A separate three-frame median appearance condition records 0.3150 mean IoU. These are preliminary temporal-localization measurements, not identity, tracking, attribute, or ownership accuracy. One video and five queries do not support broad generalization. The work demonstrates an integrated prototype and identifies the annotations and evaluation needed for stronger conclusions.')
p(t,'Keywords: temporal moment retrieval; Moment-DETR; person detection; image-text similarity; evidence fusion; reranking; CCTV.')
pg(t);hd(t,'Contents');p(t,'Update the automatic table of contents in Word after editing.');q=t.add_paragraph();f=OxmlElement('w:fldSimple');f.set(qn('w:instr'),'TOC \\o "1-3" \\h \\z \\u');q._p.append(f)
pg(t);hd(t,'List of Figures');bullets(t,['Figure 4.1 System architecture','Figure 4.2 Baseline pipeline','Figure 4.3 Method D fusion','Figure 4.4 Data flow','Figure 7.1 Mean IoU','Figure 7.2 Recall@1','Figure 7.3 Actual diagnostic frames','Figure 8.1 UI screenshot placeholder']);h(t,'List of Tables');bullets(t,['Table 4.1 Module inventory','Table 4.2 Configuration','Table 6.1 Pilot data','Table 7.1 Metrics','Table 7.2 Per-query IoU','Table 9.1 Capability boundary'])
pg(t);hd(t,'List of Abbreviations');tbl(t,['Term','Meaning'],[['CLIP','Contrastive Language–Image Pretraining'],['CCTV','Closed-Circuit Television'],['FPS','Frames per second'],['IoU','Intersection over Union'],['MR','Moment retrieval'],['NMS','Non-maximum suppression'],['QVHighlights','Query-based Video Highlights'],['R@1','Recall at rank one'],['TEF','Temporal endpoint feature'],['UI','User interface'],['YOLO','You Only Look Once']])
chapters=[
('Introduction',[('1.1 Background',['Cameras record activities, yet locating an event in a long recording can require replaying much of it. Natural-language retrieval returns candidate times for a text description. Temporal localization estimates when content occurs; it does not itself identify a person.','This project applies pretrained Moment-DETR to a video/query and optionally reranks its fixed candidates using detected-person crop evidence. Results are for human review.']),('1.2 Problem and motivation',['The task is to integrate auxiliary appearance evidence into candidate ranking and report only what the saved pilot measures. A segment matching a description is not proof of who appears in it. No biometric identity recognition or persistent identity matching is implemented.','The evaluation covers five manually defined queries from one MOT17-04-SDP video; findings are preliminary.']),('1.3 Objectives',[[ 'Inspect temporal inference and feature extraction.','Integrate a video-query application with optional person/object cues.','Recover rank fusion and missing-evidence behavior.','Verify saved pilot metrics and denominator.','List evidence and placeholders needed for submission.']]),('1.4 Scope',['The checkpoint uses two-second segments and caps input at 75 segments (about 150 seconds). The evaluation JSON contains five query intervals. The available evidence does not show MOT17 as temporal-localization training data. The checkpoint is used for inference; no project retraining is reported. YOLO11n and OpenCLIP ViT-B/32 provide optional evidence.']),('1.5 Research questions and contribution',['RQ1: Can the model be integrated into a UI returning ranked intervals? RQ2: How do baseline and Method D compare on the saved pilot? RQ3: What evidence is missing for identity, ownership, and generalization claims?','The contribution is an inference-time reranking implementation and pilot analysis, not a new architecture or trained model.']),('1.6 Organization',['Chapters 2–10 cover literature, formulation, method, implementation, methodology, results, interface, limitations/ethics, and conclusion.'])]),
('Literature Review',[('2.1 Temporal localization',['Video retrieval ranks whole recordings; moment retrieval returns intervals. Semantic relevance and boundary accuracy both matter, while proposal coverage limits reranking. Lei, Berg, and Bansal introduced QVHighlights and Moment-DETR [1].']),('2.2 Transformer set prediction',['DETR uses a transformer encoder-decoder and bipartite matching for direct object set prediction [2]. Moment-DETR adapts set prediction to temporal moments and saliency [1]. This project uses a released checkpoint and does not claim upstream benchmark results as its own.']),('2.3 Vision-language similarity',['CLIP learns transferable image representations from language supervision [3]. OpenCLIP here computes crop/text cosine similarity. This is an uncalibrated ranking signal and cannot certify a person, attribute, or action.']),('2.4 Detection, tracking, fusion',['YOLO provides class boxes and confidence. This code detects people and selected object categories but does not link boxes across frames. Fusion uses within-candidate ranks because raw model scores are not calibrated to one another.']),('2.5 Literature comparison',[('table',['Work','Approach','Strength','Limit / relevance'],[['QVHighlights/Moment-DETR [1]','Temporal moments and highlights','Baseline proposals','No identity evidence'],['DETR [2]','Transformer set prediction','Architectural lineage','Image object detection'],['CLIP [3]','Image-text representation','Zero-shot similarity','Not calibrated proof'],['YOLO','Image object detector','Person/category boxes','No ownership'],['Method D','Inference-time rank fusion','Fixed model and intervals','Single-video pilot']])]),('2.6 Research gap',['The narrow question is whether auxiliary person-appearance evidence can reorder a pretrained temporal proposal set. The pilot cannot establish general efficacy or novelty merely from combining components.']),('2.7 Metrics',['Temporal IoU is intersection duration divided by union duration. Mean top-1 IoU averages each query’s rank-one maximum overlap. Recall@1 at threshold τ is the fraction with IoU≥τ; multiple annotations use maximum overlap.'])]),
('Problem Formulation and Objectives',[('3.1 Input/output',['Let V be a video of duration T, q a text query, and cᵢ=(sᵢ,eᵢ) candidates with temporal score tᵢ. Output is a ranking of the same intervals and optional extracted clips.']),('3.2 Representations',['The wrapper samples every two seconds, uses 224-pixel center crops and CLIP ViT-B/32, normalizes video/text features, and appends two temporal endpoint features. Checkpoint metadata: 512-D visual/text features; hidden dimension 256; two encoder and decoder layers; eight heads; feedforward 1024; ten moment queries.']),('3.3 Fusion equation',['For candidate i let Kᵢ be available finite components and uᵢₖ their descending tied-rank percentile among candidates. Fᵢ = (1/|Kᵢ|)Σₖ∈Kᵢuᵢₖ when Kᵢ is nonempty; otherwise fusion is undefined and temporal order is fallback. Ties use average rank. Temporal confidence breaks fusion ties. No weights are learned. Saved D includes temporal, appearance, person-object phrase, and category components.']),('3.4 Evidence semantics',['Current appearance score max-pools crop scores across frames. Separate E takes median of per-frame maxima, requiring two valid frames. Crop boxes are not associated across frames. Object category detections do not prove ownership.']),('3.5 Criteria',['RQ1: implementation and existing records, not a usability study. RQ2: same five queries/proposals. RQ3: code has no identity or ownership metric and no held-out video.'])]),
('Existing System and Proposed Method',[('4.1 Baseline',['The wrapper samples at 0.5 Hz, normalizes CLIP ViT-B/32 features, appends normalized temporal endpoints, encodes text, and runs the checkpoint. Softmax foreground scores rank ten spans; normalized center-width spans become seconds. Saliency outputs exist, but the pilot measures moment retrieval.']),('4.2 Configuration',[('table',['Setting','Value'],[['Clip length / maximum','2 s / 75 (≈150 s)'],['Feature dimensions','512 visual and text'],['Hidden dimension','256'],['Encoder / decoder','2 / 2 layers'],['Heads / feedforward','8 / 1024'],['Moment queries','10']])]),('4.3 What changes and remains',['Method D leaves weights, architecture, and proposal intervals unchanged; it only reorders. The parser uses a limited rule vocabulary. Person crops and object phrases are scored with OpenCLIP. Still-frame action similarity is not motion recognition.']),('4.4 Missing evidence',['Available component values become descending tied-rank percentiles and are averaged without weights. Missing or unsupported evidence is omitted, not zero. Repeated object condition requires positive YOLO category detection in two of three frames; non-repetition is unknown.']),('4.5 Diagrams',[('img','fig01_architecture.png','Figure 4.1. Architecture from inspected modules.'),('img','fig02_baseline.png','Figure 4.2. Baseline inference pipeline.'),('img','fig03_method_d.png','Figure 4.3. Method D; no identity/ownership inference.'),('img','fig04_dataflow.png','Figure 4.4. Video-to-ranked-moment flow.')]),('4.6 Level of aggregation',['Frame/person-crop evidence is reduced to candidate-level scores before fusion. Different people may supply maxima across frames. Final ranking concerns intervals, not individuals.'])]),
('System Architecture and Implementation',[('5.1 Runtime',['requirements.txt pins PyTorch 2.13.0, torchvision 0.28.0, NumPy 2.2.6; Streamlit is >=1.32. OpenCLIP and Ultralytics are imported but absent from visible root requirements. FFmpeg binaries/model assets are needed. Exact pilot hardware and runtime versions are unrecorded.']),('5.2 Module inventory',[('table',['Module','Role','Limit'],[['app.py','Streamlit upload/query/results','Browser screenshot missing'],['extension/query_logic.py','Rule parsing and rank fusion','Bounded vocabulary'],['extension/attribute_matching.py','YOLO crops and OpenCLIP scores','Uncalibrated evidence'],['extension/full_pipeline.py','Compose inference and scoring','Inference only'],['extension/temporal_evidence.py','Multi-frame aggregation','No identity links'],['extension/ui_person_visualization.py','Display person evidence','Not tracking'],['run_on_video/run.py','Feature extraction and moments','75-segment limit'],['run_on_video/clip_extractor.py','Validate/extract clips','FFmpeg dependent'],['standalone_eval/run_temporal_consensus_ablation.py','Pilot comparison','Not rerun for thesis']])]),('5.3 Preprocessing',['OpenCV samples interval timestamps and converts BGR crops to RGB; crops under 16 px are skipped. OpenCLIP ViT-B/32 with OpenAI weights computes normalized cosine similarity. YOLO11n initializes when requested.']),('5.4 UI and errors',['The app accepts MP4/MOV/AVI/MKV/WEBM, text, ranking mode, threshold, and padding; session state stores results. It validates boundaries and extracts H.264/AAC clips through FFmpeg. Existing report records rendered controls but not a completed browser upload-analysis sequence.']),('5.5 Screenshots',['No authentic running-app screenshot was found. Figure 8.1 is reserved. Capture initial form, completed baseline-versus-D result with timestamps/status, and expanded clip/crop evidence.'])]),
('Experimental Setup and Methodology',[('6.1 Data boundary',[('table',['Item','Evidence','Limit'],[['Video','MOT17-04-SDP path in saved JSON','External file; licence to confirm'],['Queries','5 in JSON and annotations','One video, pilot only'],['Labels','seconds intervals in local JSON','Annotation protocol unknown'],['Proposals','10 per query','Fixed across conditions'],['Split','None','No holdout'],['Training','No training performed','MOT17 not claimed as localization training']])]),('6.2 Conditions',['A temporal-only; B single-frame appearance; D saved person-object fusion; E median three-frame appearance; F E plus repeated objects; G B plus repeated objects. Same five queries and unchanged proposals. Report says no training and no weight tuning. D and E are different conditions.']),('6.3 Procedure',['Rank-one candidate is compared to all ground-truth intervals; maximum IoU is used. Mean is across five values. Recall@1 is threshold count divided by five. Proposal oracle is a ceiling diagnostic.']),('6.4 Reproducibility',['Saved experiment uses cached CLIP crop scores and prior proposals. Three frames and two-positive-frame object quorum; 37.2-second YOLO/reranking time excludes Moment-DETR and CLIP reruns. Hardware and exact versions are absent. No fresh experiment was run here.'])]),
('Results and Discussion',[('7.1 Metrics',[('table',['Condition n=5','R@.3','R@.5','R@.7','Mean IoU'],[[labs[i],f"{M[c]['recall_at_1_iou_0_3']:.3f}",f"{M[c]['recall_at_1_iou_0_5']:.3f}",f"{M[c]['recall_at_1_iou_0_7']:.3f}",f"{M[c]['mean_top1_iou']:.4f}"] for i,c in enumerate(cs)])]),('7.2 Plots',[('img','fig05_mean_iou.png','Figure 7.1. Saved mean top-1 IoU, five queries per condition.'),('img','fig06_recall.png','Figure 7.2. Recall@1 at IoU≥0.5; R@0.7 is zero for all.')]),('7.3 Interpretation',['Baseline is 0.1940 mean IoU; D is 0.4020, +0.2080 absolute. R@.3 changes .4→.8; R@.5 .2→.4; R@.7 remains zero. These are saved pilot values, not statistical/general claims. E is separate and scores .3150. F equals E and G equals B because repeated positive object evidence did not alter rankings.']),('7.4 Per-query',[('table',['Query','A','D','E'],[['q1 red backpack','0.0000','0.5712','0.0893'],['q2 gray backpack + yellow bag','0.0000','0.4691','0.5158'],['q3 green shopping bag','0.0000','0.0000','0.0000'],['q4 white shirt walking','0.5338','0.5338','0.5338'],['q5 pink shirt walking','0.4360','0.4360','0.4360']])]),('7.5 Qualitative analysis',[('img','fig07_actual_frames.png','Figure 7.3. Actual saved MOT17-04 diagnostic frames.'),'q3 is a brief missed event; shopping bag is unsupported by the configured detector. q4/q5 rankings do not change. q1 appearance and red-backpack evidence select differing frame-local crops, not persistent people. A missing q2 detection is unknown, not proof of absence. Proposal oracle is mean max IoU .5175, R@.3 1.0, R@.5 .6, R@.7 0.0; it bounds this fixed candidate set.']),('7.6 Validity',['Five queries from one video do not justify significance or generalization. The annotation procedure is unknown. Proposal coverage caps reranking. Occlusion, lighting, crowding, resolution, crop boundaries, and language affect detector/similarity evidence. D and E differ in component setup.'])]),
('Explainability and User Interface',[('8.1 Interaction',['User selects a supported video, query, ranking mode and options; session state retains predictions and evidence. Ranked valid intervals can be played as clips. Visualization code does not prove identity.']),('8.2 Score interpretation',['Temporal confidence is Moment-DETR foreground softmax. OpenCLIP is image-text cosine similarity. YOLO confidence is detector confidence. IoU is an offline ground-truth metric. They have different scales. UI can show components but not causal explanations, stable identity, ownership, or reliable action understanding.']),('8.3 Actual screenshots',['None was found among reviewed artifacts. Figure 8.1 placeholder: capture real initial form, completed baseline/Method D result, and expanded clip/crop view. Record video/query/app version and privacy review.'])]),
('Limitations, Ethics, and Future Work',[('9.1 Technical limitations',[[ 'Five queries, one video; no significance or generalization.','CCTV domain shift is not evaluated.','Occlusion, crowding, lighting and low resolution cause missed detections.','Reranking cannot recover absent proposals.','Image-text similarity is not action recognition.','No tracking, identity, or object-ownership metric.','Input capped at roughly 150 seconds.','Saved runtime omits fresh inference.']]),('9.2 Privacy',['Footage can identify individuals. Deployment requires authorized purpose, lawful access, suitable consent/notice, minimization, retention limits, access control, and human review. False appearance matches may cause harm. Confirm video permission and ethics requirements before publication.']),('9.3 Status and future',[('table',['Status','Capability'],[['Implemented','Pretrained model, UI, crop scoring, fusion, validation, clip extraction'],['Evaluated','Five queries, one video, saved conditions'],['Unverified','Identity, tracking, ownership, generalization, significance'],['Future','Multi-video labels; person IDs/boxes; ownership labels; grouped splits; confidence intervals; stronger proposals; runtime/ethics review']])]),('9.4 Next study',['Collect permitted multi-video samples; annotate intervals, target-person IDs/boxes, and person-object relations. Split by video, prespecify fusion, and measure detector, association, ownership, and retrieval separately.'])]),
('Conclusion',[('10.1 Summary',['The system connects pretrained Moment-DETR to an interface and optional inference-time evidence reranking. YOLO/OpenCLIP supply cues; rank fusion reorders fixed intervals; validation and FFmpeg support human review. No model weights or architecture were changed.']),('10.2 Finding',['Saved pilot baseline mean IoU is .1940 and D .4020 (+.2080); R@.7 is zero. Separate median appearance E is .3150. These describe a five-query single-video pilot only. Complete personal details, data provenance, real screenshots, runtime facts, and department format before submission.'])])]
for num,(name,parts) in enumerate(chapters,1):
 pg(t);hd(t,f'Chapter {num}: {name}')
 for head,body in parts:
  h(t,head)
  for x in body:
   if isinstance(x,tuple) and x[0]=='img':pic(t,x[1],x[2])
   elif isinstance(x,tuple) and x[0]=='table':tbl(t,x[1],x[2])
   elif isinstance(x,list):bullets(t,x)
   else:p(t,x)
pg(t);hd(t,'References')
for x in ['[1] J. Lei, T. L. Berg, and M. Bansal, “QVHighlights: Detecting Moments and Highlights in Videos via Natural Language Queries,” NeurIPS 34, 2021. https://papers.neurips.cc/paper_files/paper/2021/hash/62e0973455fd26eb03e91d5741a4a3bb-Abstract.html','[2] N. Carion et al., “End-to-End Object Detection with Transformers,” ECCV, 2020, pp. 213–229. https://doi.org/10.1007/978-3-030-58452-8_13','[3] A. Radford et al., “Learning Transferable Visual Models From Natural Language Supervision,” ICML, PMLR 139, 2021, pp. 8748–8763. https://proceedings.mlr.press/v139/radford21a.html','[4] Moment-DETR code: https://github.com/jayleicn/moment_detr (accessed 24 October 2026).','[5] M. Cherti et al., “Reproducible scaling laws for contrastive language-image learning,” CVPR, 2023; OpenCLIP: https://github.com/mlfoundations/open_clip.','[6] Ultralytics YOLO documentation, https://docs.ultralytics.com/ (accessed 24 October 2026; pilot package version unknown).','[7] PyTorch documentation, https://pytorch.org/docs/. [8] FFmpeg documentation, https://ffmpeg.org/documentation.html.','[9] A. Bewley et al., “Simple Online and Realtime Tracking,” IEEE ICIP, 2016, pp. 3464–3468, doi:10.1109/ICIP.2016.7533003. Tracking background only.']:p(t,x)
pg(t);hd(t,'Appendices');h(t,'A. Reproduction commands');p(t,'Existing report documents staged use of cached data and new output folders. Verify paths; commands were not executed for this thesis.')
for x in [r'.\venv\Scripts\python.exe standalone_eval\run_automated_visual_diagnostics.py --source-results outputs\proposal_diagnosis_mot17_04_verified_20261023\results.json --annotations evaluation_data\MOT17_04_SINGLE_VIDEO_GROUND_TRUTH.json --video D:\yolo_enhance\video\MOT17-04-SDP-raw.webm --output-dir outputs\automated_visual_diagnostics_repro_20261023 --frames 3',r'.\venv\Scripts\python.exe standalone_eval\run_temporal_consensus_ablation.py --source-results outputs\proposal_diagnosis_mot17_04_verified_20261023\results.json --cached-scores outputs\automated_visual_diagnostics_repro_20261023\results.json --video D:\yolo_enhance\video\MOT17-04-SDP-raw.webm --output-dir outputs\attribute_aware_consensus_repro_20261023 --frames 3 --repeat-positive-frames 2',r'$env:PYTHONPATH="."; streamlit run app.py']:
 q=t.add_paragraph(x);q.runs[0].font.name='Consolas';q.runs[0].font.size=Pt(8)
h(t,'B. Annotation format and pseudocode');p(t,'Annotation JSON maps a query to ground_truth_intervals_seconds pairs [start,end]. Authoring and adjudication protocol is not documented. Proposals = MomentDETR(video, query); score requested cues; rank each component; average available percentiles; validate and extract.')
h(t,'C. Items before submission');bullets(t,['Student/guide/institution and official forms','Video licence and manual annotation provenance','Hardware and exact package versions','Authentic Streamlit screenshots','University layout, citation, pagination and binding rules'])
t.save(O/'Final_Year_Project_Thesis_DRAFT.docx')
# short viva companion
v=Document();setup(v,'Project Architecture and Results');hd(v,'Project Architecture and Results');p(v,'Viva reference based on the repository and saved five-query pilot.')
h(v,'One-minute explanation');p(v,'Pretrained Moment-DETR proposes ten intervals for a video/query. Optional YOLO person crops and OpenCLIP similarities become candidate evidence. Within-set tied-rank percentiles are averaged to rerank fixed proposals. No retraining, identity tracking, or ownership inference.')
pic(v,'fig01_architecture.png','Figure 1. System architecture.')
h(v,'Method D');p(v,'Fᵢ is the unweighted mean of available within-candidate tied-rank percentiles; missing is omitted. Temporal confidence breaks ties. D saved components: temporal, appearance, person-object phrase, category.')
pic(v,'fig03_method_d.png','Figure 2. Method D.')
h(v,'Pilot metrics');tbl(v,['Condition','Mean IoU','R@.3','R@.5','R@.7'],[[labs[i],f"{M[c]['mean_top1_iou']:.4f}",f"{M[c]['recall_at_1_iou_0_3']:.1f}",f"{M[c]['recall_at_1_iou_0_5']:.1f}",f"{M[c]['recall_at_1_iou_0_7']:.1f}"] for i,c in enumerate(cs)])
p(v,'Baseline-to-D mean IoU difference is +0.2080. Five queries from one video; preliminary only. E is distinct (0.3150 mean IoU).')
h(v,'Viva questions and concise answers')
for a,b in [('Is D a newly trained model?','No, it reranks fixed proposals at inference.'),('Does it identify people?','No, no tracking or identity evaluation.'),('Does an object box prove carrying?','No, category detection is not ownership.'),('What if evidence is missing?','Omit it from the available rank mean.'),('Why no broad claim?','One video, five queries, no held-out set.'),('What is next?','Multi-video, person-grounded labels and grouped evaluation.')]:
 q=v.add_paragraph();q.add_run(a+' ').bold=True;q.add_run(b)
h(v,'Limitations');bullets(v,['Proposal coverage caps ranking.','Similarity is not action recognition or calibrated truth.','No identity or ownership metric.','Screenshots and environment facts remain to supply.'])
v.save(O/'Project_Architecture_and_Results.docx')
(O/'THESIS_VERIFICATION_REPORT.md').write_text("""# Thesis Verification Report

Generated 24 October 2026. Repository D:\\jegan_repo\\FYP\\moment-detr-baseline. Existing working-tree changes preserved; source and model files not changed.

## Files inspected
app.py; extension/attribute_matching.py, full_pipeline.py, query_logic.py, temporal_evidence.py, ui_person_visualization.py; run_on_video/run.py, clip_extractor.py, model_utils.py, data_utils.py, moment_detr_ckpt/opt.json; moment_detr/config.py, model.py, transformer.py; standalone_eval/run_temporal_consensus_ablation.py; outputs/attribute_aware_consensus_20261023 results.json, report, CSV and config; verified proposal/diagnostic JSON, HTML and images; evaluation_data/MOT17_04_SINGLE_VIDEO_GROUND_TRUTH.json; README, requirements, data docs/license, existing thesis draft and DOCX; contribution, audit, single-video, proposal, soft-biometric, action/attribute, end-to-end, missing-video and dataset-plan reports.

## Verified implementation
2-second sampling, max 75 segments (~150 s), CLIP ViT-B/32, feature dimensions 512, hidden 256, two encoder/decoder layers, eight heads, ten moment queries. Method D is inference-time mean of available within-set tied-rank percentiles. Saved D uses temporal, appearance, person-object phrase, and category evidence; model and intervals remain fixed. YOLO/OpenCLIP do not track identity or prove object ownership. E is a separate median multi-frame condition.

## Metrics checked against saved JSON and report
| Condition | Mean IoU | R@.3 | R@.5 | R@.7 |
|---|---:|---:|---:|---:|
| A | .1940 | .400 | .200 | .000 |
| B | .2118 | .400 | .200 | .000 |
| D | .4020 | .800 | .400 | .000 |
| E | .3150 | .600 | .400 | .000 |
| F | .3150 | .600 | .400 | .000 |
| G | .2118 | .400 | .200 | .000 |

Each denominator is five; D-A mean difference +.2080. Saved result says proposals unchanged, no model training/weight tuning. This verifies saved figures, not a fresh replay. Video is external at D:\\yolo_enhance\\video\\MOT17-04-SDP-raw.webm.

## References checked
Metadata checked against NeurIPS, ECCV, and PMLR primary proceedings for Moment-DETR/QVHighlights, DETR, and CLIP; official software links included. Pilot's exact Ultralytics/OpenCLIP versions were not preserved.

## Unverified or missing
Student/supervisor/institution and official forms; video license and annotation protocol; hardware and pilot software versions; authentic UI screenshots for initial form, complete comparison, expanded result; no fresh experiment or source tests were run; university formatting requirements.

## Created
Final_Year_Project_Thesis_DRAFT.docx; Project_Architecture_and_Results.docx; figures/ with four code-based diagrams, two metric plots, and real saved diagnostic frames; this report; submission checklist. PDF is produced by the document renderer if available. Update Word TOC after final edits.
""",encoding='utf8')
(O/'SUBMISSION_CHECKLIST.md').write_text("""# Submission Checklist

## Student and university
- [ ] Student names, registration numbers and degree wording
- [ ] Supervisor, department, institution, city and date
- [ ] Approved certificate and declaration; signatures and seal
- [ ] Confirm acknowledgements

## Dataset and experiment
- [ ] Confirm MOT17 video source, permission/licence and frame reproduction rights
- [ ] Document who created five query annotations and the annotation/review protocol
- [ ] Record hardware, OS and exact Python/PyTorch/OpenCLIP/Ultralytics/FFmpeg versions
- [ ] Verify model and external video paths on presentation machine
- [ ] Decide if fresh reproducibility run is required; this draft checks saved results only

## Authentic Streamlit screenshots
- [ ] Actual initial upload/query controls
- [ ] Completed run showing query, baseline and Method D ranks, timestamps and evidence status
- [ ] Expanded actual clip playback and crop/overlay; conduct privacy review/redaction
- [ ] Record test video/query/app version and replace Figure 8.1 placeholder

## Final formatting
- [ ] Confirm margins, font, spacing, citation style, page numbering, page limit, binding and PDF rules
- [ ] Update Word TOC and fields after edits
- [ ] Proof citations, captions, names, dates and references
- [ ] Confirm privacy/ethics requirements
- [ ] Review final PDF and open DOCX in target Word version
""",encoding='utf8')
print(O)

